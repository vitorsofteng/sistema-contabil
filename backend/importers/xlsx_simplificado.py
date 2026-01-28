#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importador de XLSX Simplificado
================================

Importa arquivos XLSX no formato simplificado com abas:
- Informações (dados da empresa)
- DRE (demonstração de resultado)
- Balanço (balanço patrimonial)
- Indicadores (opcional)
"""

import re
from typing import Dict, Optional, List, Any
from datetime import datetime
from dataclasses import dataclass, field, asdict


@dataclass
class DadosImportados:
    """Dados extraídos do arquivo XLSX."""
    # Identificação
    ano: int = 0
    mes: int = 0
    empresa_cnpj: str = ""
    empresa_nome: str = ""
    
    # DRE
    receita_bruta: float = 0
    deducoes_receita: float = 0
    receita_liquida: float = 0
    custo_produtos_vendidos: float = 0
    lucro_bruto: float = 0
    despesas_pessoal: float = 0
    despesas_administrativas: float = 0
    despesas_comerciais: float = 0
    despesas_financeiras: float = 0
    outras_despesas: float = 0
    lucro_operacional: float = 0
    lucro_liquido: float = 0
    
    # Balanço - Ativo
    ativo_total: float = 0
    ativo_circulante: float = 0
    disponibilidades: float = 0
    caixa: float = 0
    bancos: float = 0
    aplicacoes: float = 0
    clientes: float = 0
    contas_receber: float = 0
    estoques: float = 0
    ativo_nao_circulante: float = 0
    imobilizado: float = 0
    
    # Balanço - Passivo
    passivo_total: float = 0
    passivo_circulante: float = 0
    fornecedores: float = 0
    emprestimos_cp: float = 0
    impostos_pagar: float = 0
    salarios_pagar: float = 0
    passivo_nao_circulante: float = 0
    emprestimos_lp: float = 0
    
    # Patrimônio Líquido
    patrimonio_liquido: float = 0
    capital_social: float = 0
    reservas: float = 0
    lucros_acumulados: float = 0
    
    # Meta
    arquivo_origem: str = ""
    importado_em: str = ""
    
    def to_dict(self) -> Dict:
        return asdict(self)
    
    def to_dados_mensais(self) -> Dict:
        """Converte para formato de dados mensais do sistema."""
        return {
            'ano': self.ano,
            'mes': self.mes,
            'receita': self.receita_bruta,
            'receita_bruta': self.receita_bruta,
            'deducoes_receita': self.deducoes_receita,
            'receita_liquida': self.receita_liquida,
            'custos': self.custo_produtos_vendidos,
            'custos_total': self.custo_produtos_vendidos,
            'despesas': self.despesas_administrativas + self.despesas_comerciais + self.outras_despesas,
            'despesas_operacionais': self.despesas_administrativas + self.despesas_comerciais + self.outras_despesas,
            'despesas_financeiras': self.despesas_financeiras,
            'folha': self.despesas_pessoal,
            'despesas_pessoal': self.despesas_pessoal,
            'impostos': self.deducoes_receita,
            'lucro_liquido': self.lucro_liquido,
            'caixa': self.caixa + self.bancos + self.aplicacoes,
            'disponibilidades': self.disponibilidades or (self.caixa + self.bancos + self.aplicacoes),
            'disponivel': self.disponibilidades or (self.caixa + self.bancos + self.aplicacoes),
            'bancos': self.bancos,
            'aplicacoes': self.aplicacoes,
            'clientes': self.clientes,
            'contas_receber': self.clientes,
            'estoques': self.estoques,
            'ativo_circulante': self.ativo_circulante,
            'ativo_nao_circulante': self.ativo_nao_circulante,
            'ativo_total': self.ativo_total,
            'imobilizado': self.imobilizado,
            'passivo_circulante': self.passivo_circulante,
            'passivo_nao_circulante': self.passivo_nao_circulante,
            'passivo_total': self.passivo_total,
            'fornecedores': self.fornecedores,
            'emprestimos_cp': self.emprestimos_cp,
            'emprestimos_lp': self.emprestimos_lp,
            'impostos_pagar': self.impostos_pagar,
            'salarios_pagar': self.salarios_pagar,
            'patrimonio_liquido': self.patrimonio_liquido,
            'capital_social': self.capital_social,
            'reservas': self.reservas,
            'lucros_acumulados': self.lucros_acumulados,
            'arquivo_origem': self.arquivo_origem,
        }


def normalizar_texto(texto: str) -> str:
    """Normaliza texto para comparação, removendo acentos."""
    if not texto:
        return ""
    import unicodedata
    # Remover acentos
    texto = unicodedata.normalize('NFD', texto)
    texto = texto.encode('ascii', 'ignore').decode('utf-8')
    # Manter apenas letras e números
    return re.sub(r'[^a-z0-9]', '', texto.lower().strip())


def extrair_numero(valor: Any) -> float:
    """Extrai número de célula, tratando strings e None."""
    if valor is None:
        return 0
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        # Remove formatação
        texto = re.sub(r'[R$\s\.]', '', valor).replace(',', '.')
        # Remove parênteses (números negativos)
        if texto.startswith('(') and texto.endswith(')'):
            texto = '-' + texto[1:-1]
        try:
            return float(texto)
        except ValueError:
            return 0
    return 0


def importar_xlsx_simplificado(caminho_arquivo: str) -> Optional[DadosImportados]:
    """
    Importa arquivo XLSX no formato simplificado.
    
    O arquivo deve ter as abas:
    - Informações (ou Info): dados da empresa e competência
    - DRE: demonstração de resultado
    - Balanço (ou Balanco): balanço patrimonial
    
    Args:
        caminho_arquivo: Caminho para o arquivo XLSX
        
    Returns:
        DadosImportados ou None se erro
    """
    try:
        import openpyxl
    except ImportError:
        print("Erro: openpyxl não instalado. Execute: pip install openpyxl")
        return None
    
    try:
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    except Exception as e:
        print(f"Erro ao abrir arquivo: {e}")
        return None
    
    dados = DadosImportados()
    dados.arquivo_origem = caminho_arquivo
    dados.importado_em = datetime.now().isoformat()
    
    # Mapear nomes de abas
    abas = {normalizar_texto(nome): nome for nome in wb.sheetnames}
    
    # === 1. ABA INFORMAÇÕES ===
    aba_info = None
    for nome_norm in ['informacoes', 'informacao', 'info', 'dados']:
        if nome_norm in abas:
            aba_info = wb[abas[nome_norm]]
            break
    
    if aba_info:
        for row in aba_info.iter_rows(min_row=1, max_row=20, max_col=3):
            celula_a = row[0].value if len(row) > 0 else None
            celula_b = row[1].value if len(row) > 1 else None
            
            if celula_a and celula_b:
                texto_a = normalizar_texto(str(celula_a))
                valor_b = str(celula_b)
                
                # Competência (MM/AAAA)
                if 'competencia' in texto_a or 'periodo' in texto_a:
                    match = re.search(r'(\d{1,2})[/\-](\d{4})', valor_b)
                    if match:
                        dados.mes = int(match.group(1))
                        dados.ano = int(match.group(2))
                
                # CNPJ
                if 'cnpj' in texto_a:
                    dados.empresa_cnpj = valor_b.strip()
                
                # Razão Social
                if 'razao' in texto_a:
                    dados.empresa_nome = valor_b.strip()
    
    # === 2. ABA DRE ===
    aba_dre = None
    for nome_norm in ['dre', 'resultado', 'demonstracao']:
        if nome_norm in abas:
            aba_dre = wb[abas[nome_norm]]
            break
    
    if aba_dre:
        # Mapeamento direto: texto da coluna A -> campo
        mapeamento_dre = {
            'receita bruta': 'receita_bruta',
            'deduções da receita': 'deducoes_receita',
            'custo dos produtos vendidos': 'custo_produtos_vendidos',
            'despesas com pessoal': 'despesas_pessoal',
            'despesas administrativas': 'despesas_administrativas',
            'despesas comerciais': 'despesas_comerciais',
            'despesas financeiras': 'despesas_financeiras',
            'outras despesas': 'outras_despesas',
        }
        
        for row in aba_dre.iter_rows(min_row=1, max_row=30, max_col=3):
            celula_a = row[0].value if len(row) > 0 else None
            celula_b = row[1].value if len(row) > 1 else None
            
            if celula_a and celula_b is not None:
                texto = str(celula_a).lower().strip()
                texto = texto.replace('(-)', '').strip()
                
                for chave, campo in mapeamento_dre.items():
                    if chave in texto:
                        valor = extrair_numero(celula_b)
                        setattr(dados, campo, abs(valor))
                        break
    
    # === 3. ABA BALANÇO ===
    aba_balanco = None
    for nome_norm in ['balanco', 'balancopatrimonial', 'bp', 'patrimonial']:
        if nome_norm in abas:
            aba_balanco = wb[abas[nome_norm]]
            break
    
    if aba_balanco:
        for row in aba_balanco.iter_rows(min_row=1, max_row=30, max_col=6):
            # Coluna A/B = Ativo
            celula_a = row[0].value if len(row) > 0 else None
            celula_b = row[1].value if len(row) > 1 else None
            
            # Coluna D/E = Passivo
            celula_d = row[3].value if len(row) > 3 else None
            celula_e = row[4].value if len(row) > 4 else None
            
            # === ATIVO (colunas A/B) ===
            if celula_a and celula_b is not None:
                texto = str(celula_a).lower().strip()
                valor = extrair_numero(celula_b)
                
                if 'caixa' in texto and 'equivalente' not in texto:
                    dados.caixa = abs(valor)
                elif 'banco' in texto:
                    dados.bancos = abs(valor)
                elif 'aplicaç' in texto or 'aplicac' in texto:
                    dados.aplicacoes = abs(valor)
                elif 'cliente' in texto or 'duplicata' in texto or 'contas a receber' in texto:
                    dados.clientes = abs(valor)
                elif 'estoque' in texto:
                    dados.estoques = abs(valor)
                elif 'total ativo circulante' in texto:
                    dados.ativo_circulante = abs(valor)
                elif 'imobilizado' in texto:
                    dados.imobilizado = abs(valor)
                elif 'total ativo não circulante' in texto or 'total ativo nao circulante' in texto:
                    dados.ativo_nao_circulante = abs(valor)
                elif 'total do ativo' in texto or texto == 'total ativo':
                    dados.ativo_total = abs(valor)
            
            # === PASSIVO (colunas D/E) ===
            if celula_d and celula_e is not None:
                texto = str(celula_d).lower().strip()
                valor = extrair_numero(celula_e)
                
                if 'fornecedor' in texto:
                    dados.fornecedores = abs(valor)
                elif 'empréstimo' in texto or 'emprestimo' in texto:
                    if 'cp' in texto or 'curto' in texto:
                        dados.emprestimos_cp = abs(valor)
                    elif 'lp' in texto or 'longo' in texto:
                        dados.emprestimos_lp = abs(valor)
                    else:
                        # Se não especificado, assume CP
                        dados.emprestimos_cp = abs(valor)
                elif 'imposto' in texto and 'pagar' in texto:
                    dados.impostos_pagar = abs(valor)
                elif 'salário' in texto or 'salario' in texto:
                    dados.salarios_pagar = abs(valor)
                elif 'total passivo circulante' in texto:
                    dados.passivo_circulante = abs(valor)
                elif 'total passivo não circulante' in texto or 'total passivo nao circulante' in texto:
                    dados.passivo_nao_circulante = abs(valor)
                elif 'capital social' in texto:
                    dados.capital_social = abs(valor)
                elif 'reserva' in texto:
                    dados.reservas = valor  # Pode ser negativo
                elif 'lucro' in texto and 'acumulado' in texto:
                    dados.lucros_acumulados = valor  # Pode ser negativo
                elif 'prejuízo' in texto or 'prejuizo' in texto:
                    dados.lucros_acumulados = valor  # Negativo
                elif 'total patrimônio' in texto or 'total patrimonio' in texto:
                    dados.patrimonio_liquido = valor  # Pode ser negativo
                elif 'total passivo + pl' in texto or 'total passivo+pl' in texto:
                    dados.passivo_total = abs(valor)
    
    # === CÁLCULOS DERIVADOS ===
    # Disponibilidades
    if dados.disponibilidades == 0:
        dados.disponibilidades = dados.caixa + dados.bancos + dados.aplicacoes
    
    # Receita líquida
    if dados.receita_liquida == 0 and dados.receita_bruta > 0:
        dados.receita_liquida = dados.receita_bruta - dados.deducoes_receita
    
    # Lucro bruto
    if dados.lucro_bruto == 0 and dados.receita_liquida > 0:
        dados.lucro_bruto = dados.receita_liquida - dados.custo_produtos_vendidos
    
    # Lucro líquido
    if dados.lucro_liquido == 0:
        total_despesas = (dados.despesas_pessoal + dados.despesas_administrativas + 
                         dados.despesas_comerciais + dados.despesas_financeiras + dados.outras_despesas)
        dados.lucro_operacional = dados.lucro_bruto - total_despesas
        dados.lucro_liquido = dados.lucro_operacional
    
    # Ativo total
    if dados.ativo_total == 0:
        dados.ativo_total = dados.ativo_circulante + dados.ativo_nao_circulante
    
    # Ativo não circulante
    if dados.ativo_nao_circulante == 0 and dados.imobilizado > 0:
        dados.ativo_nao_circulante = dados.imobilizado
    
    wb.close()
    return dados


def importar_multiplos_xlsx(caminhos: List[str]) -> List[DadosImportados]:
    """Importa múltiplos arquivos XLSX."""
    resultados = []
    for caminho in caminhos:
        dados = importar_xlsx_simplificado(caminho)
        if dados:
            resultados.append(dados)
    return resultados


# === TESTE ===
if __name__ == '__main__':
    import sys
    import os
    
    # Testar com os arquivos gerados
    pasta = '/mnt/user-data/outputs/empresa_problemas'
    
    if os.path.exists(pasta):
        arquivos = sorted([f for f in os.listdir(pasta) if f.endswith('.xlsx')])
        print(f"Encontrados {len(arquivos)} arquivos XLSX\n")
        
        for arquivo in arquivos:
            caminho = os.path.join(pasta, arquivo)
            print(f"Importando: {arquivo}")
            
            dados = importar_xlsx_simplificado(caminho)
            if dados:
                print(f"  Período: {dados.mes:02d}/{dados.ano}")
                print(f"  Receita Bruta: R$ {dados.receita_bruta:,.2f}")
                print(f"  Lucro Líquido: R$ {dados.lucro_liquido:,.2f}")
                print(f"  Ativo Circulante: R$ {dados.ativo_circulante:,.2f}")
                print(f"  Passivo Circulante: R$ {dados.passivo_circulante:,.2f}")
                if dados.passivo_circulante > 0:
                    liq_corr = dados.ativo_circulante / dados.passivo_circulante
                    print(f"  Liquidez Corrente: {liq_corr:.2f}")
                print(f"  Patrimônio Líquido: R$ {dados.patrimonio_liquido:,.2f}")
                print()
            else:
                print("  ERRO na importação")
    else:
        print(f"Pasta não encontrada: {pasta}")
