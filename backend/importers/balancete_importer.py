"""
Importador Inteligente de Balancetes
Sistema Contábil - Sprint 5

Baseado em balancete real de empresa brasileira.
Extrai automaticamente todos os dados financeiros do balancete.
"""

import re
import os
import json
import hashlib
from datetime import datetime, date
from typing import Dict, List, Optional, Tuple, Any
from dataclasses import dataclass, field, asdict
from decimal import Decimal
from enum import Enum


# ============================================================================
# ENUMS E CONSTANTES
# ============================================================================

class TipoConta(Enum):
    """Classificação de contas contábeis."""
    ATIVO = "ativo"
    ATIVO_CIRCULANTE = "ativo_circulante"
    ATIVO_NAO_CIRCULANTE = "ativo_nao_circulante"
    PASSIVO = "passivo"
    PASSIVO_CIRCULANTE = "passivo_circulante"
    PASSIVO_NAO_CIRCULANTE = "passivo_nao_circulante"
    PATRIMONIO_LIQUIDO = "patrimonio_liquido"
    RECEITA = "receita"
    DESPESA = "despesa"
    CUSTO = "custo"
    RESULTADO = "resultado"
    DEDUCAO = "deducao"
    APURACAO = "apuracao"


class NaturezaSaldo(Enum):
    """Natureza do saldo contábil."""
    DEVEDOR = "D"
    CREDOR = "C"


# ============================================================================
# ESTRUTURAS DE DADOS
# ============================================================================

@dataclass
class ContaContabil:
    """Representa uma conta do plano de contas."""
    codigo: str
    descricao: str
    saldo_anterior: float = 0.0
    debito: float = 0.0
    credito: float = 0.0
    saldo_atual: float = 0.0
    natureza: str = ""  # D ou C
    tipo: TipoConta = TipoConta.ATIVO
    nivel: int = 1
    conta_pai: Optional[str] = None
    
    def to_dict(self) -> Dict:
        return {
            "codigo": self.codigo,
            "descricao": self.descricao,
            "saldo_anterior": self.saldo_anterior,
            "debito": self.debito,
            "credito": self.credito,
            "saldo_atual": self.saldo_atual,
            "natureza": self.natureza,
            "tipo": self.tipo.value,
            "nivel": self.nivel
        }


@dataclass
class DadosEmpresa:
    """Dados identificadores da empresa."""
    nome: str
    cnpj: str
    cnpj_formatado: str = ""
    contador_nome: Optional[str] = None
    contador_crc: Optional[str] = None
    contador_cpf: Optional[str] = None
    socios: List[Dict] = field(default_factory=list)
    
    def __post_init__(self):
        # Limpar e formatar CNPJ
        self.cnpj = re.sub(r'\D', '', self.cnpj)
        if len(self.cnpj) == 14:
            self.cnpj_formatado = f"{self.cnpj[:2]}.{self.cnpj[2:5]}.{self.cnpj[5:8]}/{self.cnpj[8:12]}-{self.cnpj[12:]}"
    
    def to_dict(self) -> Dict:
        return {
            "nome": self.nome,
            "cnpj": self.cnpj,
            "cnpj_formatado": self.cnpj_formatado,
            "contador_nome": self.contador_nome,
            "contador_crc": self.contador_crc,
            "contador_cpf": self.contador_cpf,
            "socios": self.socios
        }


@dataclass
class ImpostosDetalhados:
    """Detalhamento de impostos."""
    iss: float = 0.0
    pis: float = 0.0
    cofins: float = 0.0
    irpj: float = 0.0
    csll: float = 0.0
    outros: float = 0.0
    
    @property
    def total(self) -> float:
        return self.iss + self.pis + self.cofins + self.irpj + self.csll + self.outros
    
    def to_dict(self) -> Dict:
        return {
            "iss": self.iss,
            "pis": self.pis,
            "cofins": self.cofins,
            "irpj": self.irpj,
            "csll": self.csll,
            "outros": self.outros,
            "total": self.total
        }


@dataclass
class DadosBalancete:
    """Estrutura completa do balancete importado."""
    # Identificação
    empresa: DadosEmpresa
    periodo_inicio: date
    periodo_fim: date
    
    # Lista de todas as contas
    contas: List[ContaContabil] = field(default_factory=list)
    
    # ========== BALANÇO PATRIMONIAL ==========
    # Ativo
    ativo_total: float = 0.0
    ativo_circulante: float = 0.0
    ativo_nao_circulante: float = 0.0
    
    # Disponibilidades
    disponivel: float = 0.0
    caixa: float = 0.0
    bancos: float = 0.0
    aplicacoes_financeiras: float = 0.0
    
    # Realizável
    clientes: float = 0.0
    duplicatas_receber: float = 0.0
    estoques: float = 0.0
    impostos_recuperar: float = 0.0
    
    # Ativo Não Circulante
    investimentos: float = 0.0
    imobilizado: float = 0.0
    intangivel: float = 0.0
    
    # Passivo
    passivo_total: float = 0.0
    passivo_circulante: float = 0.0
    passivo_nao_circulante: float = 0.0
    
    # Obrigações
    fornecedores: float = 0.0
    obrigacoes_tributarias: float = 0.0
    obrigacoes_trabalhistas: float = 0.0
    emprestimos_curto_prazo: float = 0.0
    emprestimos_longo_prazo: float = 0.0
    dividendos_pagar: float = 0.0
    
    # Impostos a Recolher (Passivo)
    impostos_recolher: ImpostosDetalhados = field(default_factory=ImpostosDetalhados)
    
    # Patrimônio Líquido
    patrimonio_liquido: float = 0.0
    capital_social: float = 0.0
    reservas_capital: float = 0.0
    reservas_lucros: float = 0.0
    lucros_acumulados: float = 0.0
    prejuizos_acumulados: float = 0.0
    
    # ========== DRE - DEMONSTRAÇÃO DO RESULTADO ==========
    receita_bruta: float = 0.0
    receita_servicos: float = 0.0
    receita_vendas: float = 0.0
    
    # Deduções da Receita
    deducoes_receita: float = 0.0
    impostos_sobre_vendas: ImpostosDetalhados = field(default_factory=ImpostosDetalhados)
    
    receita_liquida: float = 0.0
    
    # Custos
    custo_produtos_vendidos: float = 0.0
    custo_servicos_prestados: float = 0.0
    custos_total: float = 0.0
    
    lucro_bruto: float = 0.0
    
    # Despesas Operacionais
    despesas_operacionais: float = 0.0
    despesas_administrativas: float = 0.0
    despesas_comerciais: float = 0.0
    despesas_financeiras: float = 0.0
    receitas_financeiras: float = 0.0
    resultado_financeiro: float = 0.0
    
    # Outras receitas/despesas
    outras_receitas: float = 0.0
    outras_despesas: float = 0.0
    receitas_nao_operacionais: float = 0.0
    despesas_nao_operacionais: float = 0.0
    
    # Resultado
    resultado_antes_ir: float = 0.0
    provisao_ir: float = 0.0
    provisao_csll: float = 0.0
    lucro_liquido: float = 0.0
    
    # ========== METADADOS ==========
    arquivo_origem: str = ""
    data_importacao: datetime = field(default_factory=datetime.now)
    hash_arquivo: str = ""
    sistema_origem: str = ""
    
    def calcular_indicadores(self) -> Dict:
        """Calcula indicadores financeiros básicos."""
        indicadores = {}
        
        # Margens
        if self.receita_bruta > 0:
            indicadores["margem_bruta"] = ((self.receita_bruta - self.custos_total) / self.receita_bruta) * 100
            indicadores["margem_liquida"] = (self.lucro_liquido / self.receita_bruta) * 100
        
        # Liquidez
        if self.passivo_circulante > 0:
            indicadores["liquidez_corrente"] = self.ativo_circulante / self.passivo_circulante
            indicadores["liquidez_seca"] = (self.ativo_circulante - self.estoques) / self.passivo_circulante
            indicadores["liquidez_imediata"] = self.disponivel / self.passivo_circulante
        
        # Endividamento
        if self.ativo_total > 0:
            indicadores["endividamento_geral"] = (self.passivo_total / self.ativo_total) * 100
        
        if self.patrimonio_liquido > 0:
            indicadores["endividamento_pl"] = (self.passivo_total / self.patrimonio_liquido) * 100
            indicadores["roe"] = (self.lucro_liquido / self.patrimonio_liquido) * 100
        
        if self.ativo_total > 0:
            indicadores["roa"] = (self.lucro_liquido / self.ativo_total) * 100
        
        return indicadores
    
    def to_dict(self) -> Dict:
        """Converte para dicionário completo."""
        return {
            "empresa": self.empresa.to_dict(),
            "periodo": {
                "inicio": self.periodo_inicio.isoformat(),
                "fim": self.periodo_fim.isoformat(),
                "mes": self.periodo_fim.month,
                "ano": self.periodo_fim.year
            },
            "balanco_patrimonial": {
                "ativo": {
                    "total": self.ativo_total,
                    "circulante": self.ativo_circulante,
                    "nao_circulante": self.ativo_nao_circulante,
                    "disponivel": self.disponivel,
                    "caixa": self.caixa,
                    "bancos": self.bancos,
                    "aplicacoes": self.aplicacoes_financeiras,
                    "clientes": self.clientes,
                    "estoques": self.estoques,
                    "imobilizado": self.imobilizado
                },
                "passivo": {
                    "total": self.passivo_total,
                    "circulante": self.passivo_circulante,
                    "nao_circulante": self.passivo_nao_circulante,
                    "fornecedores": self.fornecedores,
                    "obrigacoes_tributarias": self.obrigacoes_tributarias,
                    "emprestimos_cp": self.emprestimos_curto_prazo,
                    "emprestimos_lp": self.emprestimos_longo_prazo,
                    "dividendos_pagar": self.dividendos_pagar,
                    "impostos_recolher": self.impostos_recolher.to_dict()
                },
                "patrimonio_liquido": {
                    "total": self.patrimonio_liquido,
                    "capital_social": self.capital_social,
                    "reservas": self.reservas_lucros,
                    "lucros_acumulados": self.lucros_acumulados
                }
            },
            "dre": {
                "receita_bruta": self.receita_bruta,
                "receita_servicos": self.receita_servicos,
                "receita_vendas": self.receita_vendas,
                "deducoes": {
                    "total": self.deducoes_receita,
                    "impostos": self.impostos_sobre_vendas.to_dict()
                },
                "receita_liquida": self.receita_liquida,
                "custos": {
                    "total": self.custos_total,
                    "cpv": self.custo_produtos_vendidos,
                    "csp": self.custo_servicos_prestados
                },
                "lucro_bruto": self.lucro_bruto,
                "despesas": {
                    "operacionais": self.despesas_operacionais,
                    "administrativas": self.despesas_administrativas,
                    "financeiras": self.despesas_financeiras
                },
                "receitas_financeiras": self.receitas_financeiras,
                "resultado_financeiro": self.resultado_financeiro,
                "outras_receitas": self.outras_receitas + self.receitas_nao_operacionais,
                "resultado_antes_ir": self.resultado_antes_ir,
                "lucro_liquido": self.lucro_liquido
            },
            "indicadores": self.calcular_indicadores(),
            "metadados": {
                "arquivo": self.arquivo_origem,
                "importado_em": self.data_importacao.isoformat(),
                "hash": self.hash_arquivo,
                "sistema": self.sistema_origem,
                "total_contas": len(self.contas)
            }
        }


@dataclass
class ResultadoImportacao:
    """Resultado da importação de balancete."""
    sucesso: bool
    mensagem: str
    dados: Optional[DadosBalancete] = None
    empresa_nova: bool = False
    empresa_id: Optional[int] = None
    avisos: List[str] = field(default_factory=list)
    erros: List[str] = field(default_factory=list)
    
    def to_dict(self) -> Dict:
        return {
            "sucesso": self.sucesso,
            "mensagem": self.mensagem,
            "empresa_nova": self.empresa_nova,
            "empresa_id": self.empresa_id,
            "avisos": self.avisos,
            "erros": self.erros,
            "dados": self.dados.to_dict() if self.dados else None
        }


# ============================================================================
# PARSER DE BALANCETE
# ============================================================================

class ParserBalancete:
    """Parser para extrair dados de texto de balancete."""
    
    # Padrões de extração
    PATTERNS = {
        # CNPJ em diversos formatos
        "cnpj": r'C\.?N\.?P\.?J\.?:?\s*(\d{2}[.\s]?\d{3}[.\s]?\d{3}[/\s]?\d{4}[-\s]?\d{2})',
        
        # Período
        "periodo": r'Per[íi]odo:?\s*(\d{2}[/.-]\d{2}[/.-]\d{4})\s*[-–a]\s*(\d{2}[/.-]\d{2}[/.-]\d{4})',
        
        # Contador
        "contador": r'Contador:?\s*\n?\s*([A-ZÀÁÂÃÉÊÍÓÔÕÚÇ][A-Za-zàáâãéêíóôõúç\s]+)',
        "crc": r'(?:CRC|Reg\.?\s*(?:no\s*)?CRC)[:\s-]*([A-Z]{2}[\s-]?\d+[/O]?\d*)',
        
        # CPF
        "cpf": r'CPF:?\s*(\d{3}[.\s]?\d{3}[.\s]?\d{3}[-\s]?\d{2})',
        
        # Sistema de origem
        "sistema": r'Sistema\s+licenciado\s+para\s+([A-ZÀÁÂÃÉÊÍÓÔÕÚÇ][A-Za-zàáâãéêíóôõúç\s]+(?:LTDA|ME|EPP)?)'
    }
    
    # Mapeamento de descrições para campos
    MAPEAMENTO_CONTAS = {
        # Ativo
        r'^ativo$': ('ativo_total', 'D'),
        r'^ativo circulante$': ('ativo_circulante', 'D'),
        r'^dispon[íi]vel$': ('disponivel', 'D'),
        r'^caixa$': ('caixa', 'D'),
        r'^caixa geral$': ('caixa', 'D'),
        r'^bancos?\s*(conta\s*movimento)?$': ('bancos', 'D'),
        r'^clientes?$': ('clientes', 'D'),
        r'^duplicatas\s*a\s*receber$': ('duplicatas_receber', 'D'),
        r'^estoque': ('estoques', 'D'),
        r'^aplica[çc][õo]es\s*financeiras?': ('aplicacoes_financeiras', 'D'),
        r'^imobilizado': ('imobilizado', 'D'),
        r'^intang[íi]vel': ('intangivel', 'D'),
        
        # Passivo
        r'^passivo$': ('passivo_total', 'C'),
        r'^passivo circulante$': ('passivo_circulante', 'C'),
        r'^passivo n[ãa]o[- ]?circulante$': ('passivo_nao_circulante', 'C'),
        r'^fornecedor': ('fornecedores', 'C'),
        r'^obriga[çc][õo]es\s*tribut[áa]rias': ('obrigacoes_tributarias', 'C'),
        r'^obriga[çc][õo]es\s*trabalhistas': ('obrigacoes_trabalhistas', 'C'),
        r'^empr[ée]stimos?$': ('emprestimos_curto_prazo', 'C'),
        r'^dividendos?\s*a\s*pagar': ('dividendos_pagar', 'C'),
        
        # Impostos a recolher
        r'^iss\s*a\s*recolher': ('impostos_recolher.iss', 'C'),
        r'^pis\s*a\s*recolher': ('impostos_recolher.pis', 'C'),
        r'^cofins\s*a\s*recolher': ('impostos_recolher.cofins', 'C'),
        r'^(?:imposto\s*de\s*renda|irpj?)\s*a\s*recolher': ('impostos_recolher.irpj', 'C'),
        r'^(?:contribui[çc][ãa]o\s*social|csll?)\s*a\s*recolher': ('impostos_recolher.csll', 'C'),
        
        # Patrimônio Líquido
        r'^patrim[ôo]nio\s*l[íi]quido$': ('patrimonio_liquido', 'C'),
        r'^capital\s*social$': ('capital_social', 'C'),
        r'^capital\s*subscrito$': ('capital_social', 'C'),
        r'^lucros?\s*(?:ou\s*preju[íi]zos?)?\s*acumulados?': ('lucros_acumulados', 'C'),
        r'^reservas?\s*de\s*lucros?': ('reservas_lucros', 'C'),
        
        # DRE - Receitas
        r'^receita\s*bruta': ('receita_bruta', 'C'),
        r'^receita.*servi[çc]os?\s*prestados?': ('receita_servicos', 'C'),
        r'^servi[çc]os\s*prestados$': ('receita_servicos', 'C'),
        r'^receita.*vendas?': ('receita_vendas', 'C'),
        
        # DRE - Deduções
        r'^\(-\)\s*dedu[çc][õo]es': ('deducoes_receita', 'D'),
        r'^\(-\)\s*iss$': ('impostos_sobre_vendas.iss', 'D'),
        r'^\(-\)\s*pis$': ('impostos_sobre_vendas.pis', 'D'),
        r'^\(-\)\s*cofins$': ('impostos_sobre_vendas.cofins', 'D'),
        r'^\(-\)\s*(?:imposto\s*de\s*renda|irpj?)': ('impostos_sobre_vendas.irpj', 'D'),
        r'^\(-\)\s*(?:contribui[çc][ãa]o\s*social|csll?)': ('impostos_sobre_vendas.csll', 'D'),
        
        # DRE - Custos
        r'^custos?\s*(?:dos?\s*)?(?:produtos?|mercadorias?)\s*vendid': ('custo_produtos_vendidos', 'D'),
        r'^custos?\s*(?:dos?\s*)?servi[çc]os?\s*(?:prestados?|vendidos?)': ('custo_servicos_prestados', 'D'),
        
        # DRE - Despesas
        r'^despesas?\s*operacionais?$': ('despesas_operacionais', 'D'),
        r'^despesas?\s*administrativas?': ('despesas_administrativas', 'D'),
        r'^despesas?\s*comerciais?': ('despesas_comerciais', 'D'),
        r'^despesas?\s*financeiras?': ('despesas_financeiras', 'D'),
        r'^juros\s*passivos': ('despesas_financeiras', 'D'),
        
        # DRE - Receitas Financeiras
        r'^receitas?\s*financeiras?': ('receitas_financeiras', 'C'),
        r'^rendimento.*aplica[çc][ãa]o': ('receitas_financeiras', 'C'),
        r'^receitas?\s*n[ãa]o\s*operacionais?': ('receitas_nao_operacionais', 'C'),
        
        # Resultado
        r'^lucro\s*(?:l[íi]quido\s*)?(?:do\s*)?exerc[íi]cio': ('lucro_liquido', 'C'),
        r'^resultado\s*l[íi]quido.*irpj': ('resultado_antes_ir', 'C'),
        r'^lucro\s*bruto': ('lucro_bruto', 'C'),
        r'^resultado\s*bruto': ('lucro_bruto', 'C'),
    }
    
    def extrair_cnpj(self, texto: str) -> Optional[str]:
        """Extrai CNPJ do texto."""
        match = re.search(self.PATTERNS["cnpj"], texto, re.IGNORECASE)
        if match:
            cnpj = re.sub(r'\D', '', match.group(1))
            if len(cnpj) == 14:
                return cnpj
        return None
    
    def extrair_periodo(self, texto: str) -> Tuple[Optional[date], Optional[date]]:
        """Extrai período do balancete."""
        match = re.search(self.PATTERNS["periodo"], texto, re.IGNORECASE)
        if match:
            try:
                data_ini = self._parse_data(match.group(1))
                data_fim = self._parse_data(match.group(2))
                return data_ini, data_fim
            except:
                pass
        return None, None
    
    def _parse_data(self, texto: str) -> date:
        """Converte string de data para objeto date."""
        texto = texto.replace('.', '/').replace('-', '/')
        partes = texto.split('/')
        dia, mes, ano = int(partes[0]), int(partes[1]), int(partes[2])
        if ano < 100:
            ano += 2000
        return date(ano, mes, dia)
    
    def extrair_nome_empresa(self, texto: str) -> Optional[str]:
        """Extrai nome da empresa."""
        # Procurar padrão "Empresa: NOME"
        patterns = [
            r'Empresa:?\s*([A-ZÀÁÂÃÉÊÍÓÔÕÚÇ][A-ZÀÁÂÃÉÊÍÓÔÕÚÇ\s\.\-&0-9]+(?:LTDA|ME|EPP|EIRELI|S/?A|CIA\s*LTDA))',
            r'^([A-ZÀÁÂÃÉÊÍÓÔÕÚÇ][A-ZÀÁÂÃÉÊÍÓÔÕÚÇ\s\.\-&0-9]+(?:LTDA|ME|EPP|EIRELI|S/?A|CIA\s*LTDA))\s*$'
        ]
        
        for pattern in patterns:
            match = re.search(pattern, texto, re.MULTILINE | re.IGNORECASE)
            if match:
                nome = match.group(1).strip()
                # Limpar
                nome = re.sub(r'\s+', ' ', nome)
                nome = re.sub(r'\s*C\.?N\.?P\.?J.*$', '', nome, flags=re.IGNORECASE)
                if len(nome) > 5:
                    return nome.upper()
        
        return None
    
    def extrair_contador(self, texto: str) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        """Extrai nome, CRC e CPF do contador."""
        nome = None
        crc = None
        cpf = None
        
        # Nome do contador (geralmente após "Contador")
        match = re.search(r'Contador[:\s]*\n?\s*([A-ZÀÁÂÃÉÊÍÓÔÕÚÇ][A-Za-zàáâãéêíóôõúç\s]+)', texto, re.IGNORECASE)
        if match:
            nome = match.group(1).strip()
            nome = re.sub(r'\s+', ' ', nome)
        
        # CRC
        match = re.search(r'(?:CRC|Reg\.?\s*(?:no\s*)?CRC)[:\s-]*([A-Z]{2}[\s-]?[\dO]+)', texto, re.IGNORECASE)
        if match:
            crc = match.group(1).strip()
        
        # Tentar outro padrão
        if not crc:
            match = re.search(r'(?:sob\s+o\s+No?\.?\s*)([A-Z]{2}[\dO]+)', texto, re.IGNORECASE)
            if match:
                crc = match.group(1).strip()
        
        # CPF do contador (geralmente logo após o nome/CRC)
        matches = re.findall(self.PATTERNS["cpf"], texto)
        if matches:
            cpf = matches[0]  # Primeiro CPF encontrado geralmente é do contador
        
        return nome, crc, cpf
    
    def extrair_sistema(self, texto: str) -> Optional[str]:
        """Extrai nome do sistema de origem."""
        match = re.search(self.PATTERNS["sistema"], texto, re.IGNORECASE)
        if match:
            return match.group(1).strip()
        return None
    
    def parse_valor_brasileiro(self, texto: str) -> float:
        """Converte valor no formato brasileiro para float."""
        if not texto or texto.strip() == '' or texto.strip() == '0':
            return 0.0
        
        texto = str(texto).strip()
        
        # Remover indicador de natureza
        texto = re.sub(r'[DC]$', '', texto).strip()
        
        # Formato brasileiro: 1.234.567,89
        if ',' in texto:
            texto = texto.replace('.', '').replace(',', '.')
        
        try:
            return abs(float(texto))
        except:
            return 0.0
    
    def identificar_conta(self, descricao: str) -> Tuple[Optional[str], Optional[str]]:
        """Identifica o campo correspondente à descrição da conta."""
        desc_lower = descricao.lower().strip()
        
        for pattern, (campo, natureza) in self.MAPEAMENTO_CONTAS.items():
            if re.match(pattern, desc_lower, re.IGNORECASE):
                return campo, natureza
        
        return None, None
    
    def classificar_tipo_conta(self, codigo: str, descricao: str) -> TipoConta:
        """Classifica o tipo da conta baseado no código e descrição."""
        desc_lower = descricao.lower()
        
        # Por código (plano de contas padrão)
        if codigo:
            primeiro = codigo[0]
            if primeiro == '1':
                return TipoConta.ATIVO
            elif primeiro == '2':
                if 'patrimônio' in desc_lower or 'patrimonio' in desc_lower:
                    return TipoConta.PATRIMONIO_LIQUIDO
                return TipoConta.PASSIVO
            elif primeiro == '3':
                return TipoConta.RECEITA
            elif primeiro == '4':
                if '(-)' in descricao:
                    return TipoConta.DEDUCAO
                return TipoConta.DESPESA
            elif primeiro == '5':
                return TipoConta.RESULTADO
        
        # Por palavras-chave
        if 'ativo' in desc_lower:
            return TipoConta.ATIVO
        elif 'passivo' in desc_lower:
            return TipoConta.PASSIVO
        elif 'patrimônio' in desc_lower or 'patrimonio' in desc_lower or 'capital' in desc_lower:
            return TipoConta.PATRIMONIO_LIQUIDO
        elif 'receita' in desc_lower or 'venda' in desc_lower:
            return TipoConta.RECEITA
        elif 'custo' in desc_lower:
            return TipoConta.CUSTO
        elif 'despesa' in desc_lower:
            return TipoConta.DESPESA
        elif '(-)' in descricao:
            return TipoConta.DEDUCAO
        
        return TipoConta.RESULTADO


# ============================================================================
# IMPORTADOR BASE
# ============================================================================

class ImportadorBalanceteBase:
    """Classe base para importadores de balancete."""
    
    def __init__(self):
        self.parser = ParserBalancete()
        self.erros = []
        self.avisos = []
    
    def _calcular_hash(self, caminho: str) -> str:
        """Calcula hash MD5 do arquivo."""
        hasher = hashlib.md5()
        with open(caminho, 'rb') as f:
            for chunk in iter(lambda: f.read(4096), b''):
                hasher.update(chunk)
        return hasher.hexdigest()
    
    def _atribuir_valor(self, dados: DadosBalancete, campo: str, valor: float):
        """Atribui valor ao campo do balancete, incluindo campos aninhados."""
        if '.' in campo:
            partes = campo.split('.')
            obj = getattr(dados, partes[0])
            setattr(obj, partes[1], valor)
        else:
            setattr(dados, campo, valor)
    
    def _processar_conta(self, dados: DadosBalancete, conta: ContaContabil):
        """Processa uma conta e atribui valores ao balancete."""
        campo, natureza = self.parser.identificar_conta(conta.descricao)
        
        if campo:
            valor = conta.saldo_atual
            self._atribuir_valor(dados, campo, valor)
        
        # Adicionar à lista de contas
        dados.contas.append(conta)
    
    def _calcular_derivados(self, dados: DadosBalancete):
        """Calcula valores derivados após importação."""
        # Disponível
        if dados.disponivel == 0:
            dados.disponivel = dados.caixa + dados.bancos + dados.aplicacoes_financeiras
        
        # Clientes
        if dados.clientes == 0 and dados.duplicatas_receber > 0:
            dados.clientes = dados.duplicatas_receber
        
        # Custos total
        dados.custos_total = dados.custo_produtos_vendidos + dados.custo_servicos_prestados
        
        # Receita de serviços
        if dados.receita_servicos == 0 and dados.receita_bruta > 0 and dados.receita_vendas == 0:
            dados.receita_servicos = dados.receita_bruta
        
        # Receita líquida
        if dados.receita_liquida == 0:
            dados.receita_liquida = dados.receita_bruta - dados.deducoes_receita
        
        # Lucro bruto
        if dados.lucro_bruto == 0:
            dados.lucro_bruto = dados.receita_liquida - dados.custos_total
        
        # Deduções da receita (se não informado)
        if dados.deducoes_receita == 0:
            dados.deducoes_receita = dados.impostos_sobre_vendas.total
        
        # Resultado financeiro
        dados.resultado_financeiro = dados.receitas_financeiras - dados.despesas_financeiras
        
        # Empréstimos
        if dados.emprestimos_longo_prazo == 0 and dados.passivo_nao_circulante > 0:
            dados.emprestimos_longo_prazo = dados.passivo_nao_circulante
        
        # Ativo não circulante
        if dados.ativo_nao_circulante == 0:
            dados.ativo_nao_circulante = dados.ativo_total - dados.ativo_circulante
        
        # Passivo não circulante  
        if dados.passivo_nao_circulante == 0 and dados.passivo_total > 0 and dados.passivo_circulante > 0:
            # Cuidado: passivo_total pode incluir PL em alguns planos de contas
            if dados.passivo_total > dados.passivo_circulante:
                calc = dados.passivo_total - dados.passivo_circulante - dados.patrimonio_liquido
                if calc > 0:
                    dados.passivo_nao_circulante = calc


# ============================================================================
# IMPORTADOR PDF
# ============================================================================

class ImportadorPDF(ImportadorBalanceteBase):
    """Importador de balancetes em formato PDF."""
    
    def importar(self, caminho: str) -> ResultadoImportacao:
        """Importa balancete de arquivo PDF."""
        try:
            texto = self._extrair_texto(caminho)
            
            if not texto:
                return ResultadoImportacao(
                    sucesso=False,
                    mensagem="Não foi possível extrair texto do PDF"
                )
            
            return self._processar_texto(texto, caminho)
            
        except Exception as e:
            return ResultadoImportacao(
                sucesso=False,
                mensagem=f"Erro ao importar PDF: {str(e)}",
                erros=[str(e)]
            )
    
    def _extrair_texto(self, caminho: str) -> str:
        """Extrai texto do PDF usando diferentes bibliotecas."""
        texto = ""
        
        # Tentar PyMuPDF (fitz)
        try:
            import fitz
            doc = fitz.open(caminho)
            for pagina in doc:
                texto += pagina.get_text()
            doc.close()
            if texto.strip():
                return texto
        except:
            pass
        
        # Tentar pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(caminho) as pdf:
                for pagina in pdf.pages:
                    texto += (pagina.extract_text() or "") + "\n"
            if texto.strip():
                return texto
        except:
            pass
        
        # Tentar PyPDF2
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(caminho)
            for pagina in reader.pages:
                texto += (pagina.extract_text() or "") + "\n"
            if texto.strip():
                return texto
        except:
            pass
        
        return texto
    
    def _processar_texto(self, texto: str, caminho: str) -> ResultadoImportacao:
        """Processa texto extraído do PDF."""
        self.avisos = []
        self.erros = []
        
        # Extrair dados básicos
        cnpj = self.parser.extrair_cnpj(texto)
        if not cnpj:
            return ResultadoImportacao(
                sucesso=False,
                mensagem="CNPJ não encontrado no documento. Verifique se é um balancete válido."
            )
        
        nome_empresa = self.parser.extrair_nome_empresa(texto)
        if not nome_empresa:
            self.avisos.append("Nome da empresa não identificado claramente")
            nome_empresa = "EMPRESA NÃO IDENTIFICADA"
        
        periodo_inicio, periodo_fim = self.parser.extrair_periodo(texto)
        if not periodo_inicio or not periodo_fim:
            self.avisos.append("Período não identificado, usando mês atual")
            periodo_fim = date.today().replace(day=1)
            periodo_inicio = periodo_fim.replace(day=1)
        
        contador_nome, contador_crc, contador_cpf = self.parser.extrair_contador(texto)
        sistema = self.parser.extrair_sistema(texto)
        
        # Criar estruturas
        empresa = DadosEmpresa(
            nome=nome_empresa,
            cnpj=cnpj,
            contador_nome=contador_nome,
            contador_crc=contador_crc,
            contador_cpf=contador_cpf
        )
        
        dados = DadosBalancete(
            empresa=empresa,
            periodo_inicio=periodo_inicio,
            periodo_fim=periodo_fim,
            arquivo_origem=os.path.basename(caminho),
            hash_arquivo=self._calcular_hash(caminho),
            sistema_origem=sistema or "Desconhecido"
        )
        
        # Extrair contas linha por linha
        self._extrair_contas_texto(texto, dados)
        
        # Calcular valores derivados
        self._calcular_derivados(dados)
        
        return ResultadoImportacao(
            sucesso=True,
            mensagem=f"Balancete importado com sucesso: {nome_empresa}",
            dados=dados,
            avisos=self.avisos,
            erros=self.erros
        )
    
    def _extrair_contas_texto(self, texto: str, dados: DadosBalancete):
        """Extrai contas contábeis do texto."""
        linhas = texto.split('\n')
        
        for linha in linhas:
            linha = linha.strip()
            if not linha or len(linha) < 5:
                continue
            
            # Padrão: código descrição valor valor valor valor [D/C]
            # Exemplo: 1 ATIVO 0,00 7.201.123,29 1.138.242,70 6.062.880,59D
            
            # Tentar diferentes padrões
            conta = self._parse_linha_conta(linha)
            if conta:
                self._processar_conta(dados, conta)
    
    def _parse_linha_conta(self, linha: str) -> Optional[ContaContabil]:
        """Tenta parsear uma linha como conta contábil."""
        # Padrão 1: código numérico + descrição + 4 valores + natureza
        match = re.match(
            r'^(\d+)\s+([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][^\d]+?)\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s*([DC])?$',
            linha,
            re.IGNORECASE
        )
        
        if match:
            return ContaContabil(
                codigo=match.group(1),
                descricao=match.group(2).strip(),
                saldo_anterior=self.parser.parse_valor_brasileiro(match.group(3)),
                debito=self.parser.parse_valor_brasileiro(match.group(4)),
                credito=self.parser.parse_valor_brasileiro(match.group(5)),
                saldo_atual=self.parser.parse_valor_brasileiro(match.group(6)),
                natureza=match.group(7) or '',
                tipo=self.parser.classificar_tipo_conta(match.group(1), match.group(2))
            )
        
        # Padrão 2: descrição duplicada (comum em alguns sistemas)
        # Exemplo: "CAIXA4 CAIXA 0,00 10.000,00 0,00 10.000,00D"
        match = re.match(
            r'^([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][^\d]+?)(\d+)\s+\1\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s*([DC])?$',
            linha,
            re.IGNORECASE
        )
        
        if match:
            return ContaContabil(
                codigo=match.group(2),
                descricao=match.group(1).strip(),
                saldo_anterior=self.parser.parse_valor_brasileiro(match.group(3)),
                debito=self.parser.parse_valor_brasileiro(match.group(4)),
                credito=self.parser.parse_valor_brasileiro(match.group(5)),
                saldo_atual=self.parser.parse_valor_brasileiro(match.group(6)),
                natureza=match.group(7) or '',
                tipo=self.parser.classificar_tipo_conta(match.group(2), match.group(1))
            )
        
        return None


# ============================================================================
# IMPORTADOR EXCEL (XLS/XLSX)
# ============================================================================

class ImportadorExcel(ImportadorBalanceteBase):
    """Importador de balancetes em formato Excel."""
    
    def importar(self, caminho: str) -> ResultadoImportacao:
        """Importa balancete de arquivo Excel."""
        try:
            dados_planilha = self._ler_planilha(caminho)
            
            if not dados_planilha:
                return ResultadoImportacao(
                    sucesso=False,
                    mensagem="Não foi possível ler o arquivo Excel"
                )
            
            return self._processar_planilha(dados_planilha, caminho)
            
        except Exception as e:
            return ResultadoImportacao(
                sucesso=False,
                mensagem=f"Erro ao importar Excel: {str(e)}",
                erros=[str(e)]
            )
    
    def _ler_planilha(self, caminho: str) -> Optional[List[List]]:
        """Lê dados da planilha Excel."""
        dados = None
        
        # Tentar openpyxl (xlsx)
        try:
            from openpyxl import load_workbook
            wb = load_workbook(caminho, data_only=True)
            ws = wb.active
            dados = []
            for row in ws.iter_rows():
                dados.append([cell.value for cell in row])
            wb.close()
            if dados:
                return dados
        except:
            pass
        
        # Tentar xlrd (xls antigo)
        try:
            import xlrd
            wb = xlrd.open_workbook(caminho)
            ws = wb.sheet_by_index(0)
            dados = []
            for row_idx in range(ws.nrows):
                dados.append([ws.cell_value(row_idx, col) for col in range(ws.ncols)])
            if dados:
                return dados
        except:
            pass
        
        # Tentar pandas como fallback
        try:
            import pandas as pd
            df = pd.read_excel(caminho, header=None)
            dados = df.values.tolist()
            if dados:
                return dados
        except:
            pass
        
        return dados
    
    def _processar_planilha(self, dados: List[List], caminho: str) -> ResultadoImportacao:
        """Processa dados da planilha."""
        self.avisos = []
        self.erros = []
        
        # Converter para texto para extrair metadados
        texto_completo = "\n".join(
            " ".join(str(c) if c else "" for c in row)
            for row in dados
        )
        
        # Extrair dados básicos
        cnpj = self.parser.extrair_cnpj(texto_completo)
        if not cnpj:
            return ResultadoImportacao(
                sucesso=False,
                mensagem="CNPJ não encontrado no documento"
            )
        
        nome_empresa = self.parser.extrair_nome_empresa(texto_completo)
        if not nome_empresa:
            self.avisos.append("Nome da empresa não identificado")
            nome_empresa = "EMPRESA NÃO IDENTIFICADA"
        
        periodo_inicio, periodo_fim = self.parser.extrair_periodo(texto_completo)
        if not periodo_inicio or not periodo_fim:
            self.avisos.append("Período não identificado")
            periodo_fim = date.today()
            periodo_inicio = periodo_fim.replace(day=1)
        
        contador_nome, contador_crc, contador_cpf = self.parser.extrair_contador(texto_completo)
        sistema = self.parser.extrair_sistema(texto_completo)
        
        # Criar estruturas
        empresa = DadosEmpresa(
            nome=nome_empresa,
            cnpj=cnpj,
            contador_nome=contador_nome,
            contador_crc=contador_crc,
            contador_cpf=contador_cpf
        )
        
        balancete = DadosBalancete(
            empresa=empresa,
            periodo_inicio=periodo_inicio,
            periodo_fim=periodo_fim,
            arquivo_origem=os.path.basename(caminho),
            hash_arquivo=self._calcular_hash(caminho),
            sistema_origem=sistema or "Desconhecido"
        )
        
        # Extrair contas das linhas
        self._extrair_contas_planilha(dados, balancete)
        
        # Calcular derivados
        self._calcular_derivados(balancete)
        
        return ResultadoImportacao(
            sucesso=True,
            mensagem=f"Balancete importado com sucesso: {nome_empresa}",
            dados=balancete,
            avisos=self.avisos,
            erros=self.erros
        )
    
    def _extrair_contas_planilha(self, dados: List[List], balancete: DadosBalancete):
        """Extrai contas da planilha."""
        for row in dados:
            if not row or len(row) < 4:
                continue
            
            # Primeira coluna geralmente é código
            codigo = str(row[0]).strip() if row[0] else ""
            
            # Pular linhas sem código numérico
            if not codigo or not re.match(r'^\d+$', codigo):
                continue
            
            # Segunda coluna é descrição
            descricao = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            
            if not descricao:
                continue
            
            # Colunas de valores (ordem padrão: saldo_ant, débito, crédito, saldo_atual)
            valores = []
            for i in range(2, min(6, len(row))):
                val = row[i] if len(row) > i else 0
                valores.append(self._parse_valor_celula(val))
            
            while len(valores) < 4:
                valores.append(0.0)
            
            # Verificar natureza na última coluna
            natureza = ''
            if len(row) > 6:
                nat = str(row[6]).strip().upper() if row[6] else ''
                if nat in ['D', 'C']:
                    natureza = nat
            
            conta = ContaContabil(
                codigo=codigo,
                descricao=descricao,
                saldo_anterior=valores[0],
                debito=valores[1],
                credito=valores[2],
                saldo_atual=valores[3],
                natureza=natureza,
                tipo=self.parser.classificar_tipo_conta(codigo, descricao)
            )
            
            self._processar_conta(balancete, conta)
    
    def _parse_valor_celula(self, valor) -> float:
        """Converte valor de célula para float."""
        if valor is None:
            return 0.0
        
        if isinstance(valor, (int, float)):
            return abs(float(valor))
        
        return self.parser.parse_valor_brasileiro(str(valor))


# ============================================================================
# IMPORTADOR CSV
# ============================================================================

class ImportadorCSV(ImportadorBalanceteBase):
    """Importador de balancetes em formato CSV."""
    
    def importar(self, caminho: str) -> ResultadoImportacao:
        """Importa balancete de arquivo CSV."""
        try:
            import csv
            
            # Detectar delimitador
            with open(caminho, 'r', encoding='utf-8-sig') as f:
                amostra = f.read(4096)
                dialect = csv.Sniffer().sniff(amostra, delimiters=';,\t')
                f.seek(0)
                reader = csv.reader(f, dialect)
                dados = list(reader)
            
            if not dados:
                return ResultadoImportacao(
                    sucesso=False,
                    mensagem="Arquivo CSV vazio"
                )
            
            # Usar processador do Excel (mesmo formato)
            importador_excel = ImportadorExcel()
            importador_excel._calcular_hash = self._calcular_hash
            return importador_excel._processar_planilha(dados, caminho)
            
        except Exception as e:
            return ResultadoImportacao(
                sucesso=False,
                mensagem=f"Erro ao importar CSV: {str(e)}",
                erros=[str(e)]
            )


# ============================================================================
# GERENCIADOR DE IMPORTAÇÃO
# ============================================================================

class GerenciadorImportacao:
    """Gerenciador principal de importação de balancetes."""
    
    def __init__(self, repositorio_empresas=None):
        self.importador_pdf = ImportadorPDF()
        self.importador_excel = ImportadorExcel()
        self.importador_csv = ImportadorCSV()
        self.repositorio = repositorio_empresas
    
    def importar(self, caminho: str) -> ResultadoImportacao:
        """
        Importa arquivo de balancete automaticamente.
        
        Fluxo:
        1. Detecta o tipo de arquivo pela extensão
        2. Usa o importador apropriado
        3. Verifica se empresa existe pelo CNPJ
        4. Cria empresa se necessário
        5. Retorna resultado completo com dados estruturados
        """
        if not os.path.exists(caminho):
            return ResultadoImportacao(
                sucesso=False,
                mensagem=f"Arquivo não encontrado: {caminho}"
            )
        
        # Detectar tipo
        extensao = os.path.splitext(caminho)[1].lower()
        
        if extensao == '.pdf':
            resultado = self.importador_pdf.importar(caminho)
        elif extensao in ['.xls', '.xlsx', '.xlsm']:
            resultado = self.importador_excel.importar(caminho)
        elif extensao == '.csv':
            resultado = self.importador_csv.importar(caminho)
        else:
            return ResultadoImportacao(
                sucesso=False,
                mensagem=f"Formato não suportado: {extensao}. Use PDF, XLS, XLSX ou CSV."
            )
        
        # Processar empresa se importação ok
        if resultado.sucesso and resultado.dados and self.repositorio:
            self._processar_empresa(resultado)
        
        return resultado
    
    def _processar_empresa(self, resultado: ResultadoImportacao):
        """Verifica/cria empresa no repositório."""
        cnpj = resultado.dados.empresa.cnpj
        
        # Buscar empresa
        empresa = self.repositorio.buscar_por_cnpj(cnpj)
        
        if empresa:
            resultado.empresa_id = empresa.id
            resultado.empresa_nova = False
            resultado.mensagem += f" | Empresa existente (ID: {empresa.id})"
        else:
            # Criar nova
            nova = self.repositorio.criar_empresa(
                nome=resultado.dados.empresa.nome,
                cnpj=cnpj
            )
            resultado.empresa_id = nova.id
            resultado.empresa_nova = True
            resultado.mensagem += f" | Nova empresa criada (ID: {nova.id})"
    
    def para_dados_mensais(self, dados: DadosBalancete) -> Dict:
        """Converte para formato de dados mensais do sistema."""
        return {
            "empresa_cnpj": dados.empresa.cnpj,
            "empresa_nome": dados.empresa.nome,
            "mes": dados.periodo_fim.month,
            "ano": dados.periodo_fim.year,
            "periodo_inicio": dados.periodo_inicio.isoformat(),
            "periodo_fim": dados.periodo_fim.isoformat(),
            
            # Balanço
            "ativo_total": dados.ativo_total,
            "ativo_circulante": dados.ativo_circulante,
            "passivo_circulante": dados.passivo_circulante,
            "patrimonio_liquido": dados.patrimonio_liquido,
            "disponivel": dados.disponivel,
            
            # DRE
            "receita_bruta": dados.receita_bruta,
            "custos": dados.custos_total,
            "despesas_operacionais": dados.despesas_operacionais,
            "despesas_financeiras": dados.despesas_financeiras,
            "lucro_liquido": dados.lucro_liquido,
            
            # Impostos
            "impostos": dados.deducoes_receita,
            
            # Metadados
            "importado_de": dados.arquivo_origem,
            "sistema_origem": dados.sistema_origem
        }


# ============================================================================
# FUNÇÕES DE CONVENIÊNCIA
# ============================================================================

def importar_balancete(caminho: str, repositorio=None) -> ResultadoImportacao:
    """Importa balancete de qualquer formato suportado."""
    gerenciador = GerenciadorImportacao(repositorio)
    return gerenciador.importar(caminho)


def importar_e_converter(caminho: str) -> Tuple[bool, Dict, str]:
    """Importa e converte para formato do sistema."""
    resultado = importar_balancete(caminho)
    
    if not resultado.sucesso:
        return False, {}, resultado.mensagem
    
    gerenciador = GerenciadorImportacao()
    dados = gerenciador.para_dados_mensais(resultado.dados)
    
    return True, dados, resultado.mensagem


# ============================================================================
# TESTE
# ============================================================================

if __name__ == "__main__":
    import sys
    
    if len(sys.argv) > 1:
        caminho = sys.argv[1]
        resultado = importar_balancete(caminho)
        
        print(f"\nSucesso: {resultado.sucesso}")
        print(f"Mensagem: {resultado.mensagem}")
        
        if resultado.avisos:
            print(f"Avisos: {resultado.avisos}")
        
        if resultado.sucesso and resultado.dados:
            print("\n" + "="*60)
            print("DADOS EXTRAÍDOS")
            print("="*60)
            dados_dict = resultado.dados.to_dict()
            print(json.dumps(dados_dict, indent=2, ensure_ascii=False, default=str))
