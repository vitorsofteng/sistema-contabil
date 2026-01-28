#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importação Avançada de Dados - Sistema Contábil
================================================

Suporte a CSV, Excel, OFX (extratos bancários) e XML NFe
"""

import os
import io
import json
import hashlib
import re
from datetime import datetime, date
from typing import Optional, List, Dict, Any, Tuple
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
import xml.etree.ElementTree as ET

from sqlalchemy import (
    Column, Integer, String, Float, Boolean, Text, DateTime,
    ForeignKey, Index, func, desc
)
from sqlalchemy.orm import relationship, Session

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from data.database import Base, get_db

# Bibliotecas opcionais
try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


# =============================================================================
# MODELOS
# =============================================================================

class Importacao(Base):
    """Modelo de importação."""
    __tablename__ = 'importacoes'
    
    id = Column(Integer, primary_key=True, autoincrement=True)
    organizacao_id = Column(Integer, ForeignKey('organizacoes.id'))
    empresa_id = Column(Integer, ForeignKey('empresas.id'))
    contador_id = Column(Integer, ForeignKey('contadores.id'), nullable=False)
    
    nome_arquivo = Column(String(255), nullable=False)
    tipo_arquivo = Column(String(20), nullable=False)
    tamanho_bytes = Column(Integer)
    hash_arquivo = Column(String(64), index=True)
    
    status = Column(String(30), default='pending')
    total_registros = Column(Integer, default=0)
    registros_importados = Column(Integer, default=0)
    registros_duplicados = Column(Integer, default=0)
    registros_erro = Column(Integer, default=0)
    
    mapeamento_id = Column(Integer)
    mapeamento_json = Column(Text)
    erros_json = Column(Text)
    resumo = Column(Text)
    
    iniciado_em = Column(DateTime)
    finalizado_em = Column(DateTime)
    created_at = Column(DateTime, default=func.now())
    
    registros = relationship("RegistroImportado", back_populates="importacao")
    
    def to_dict(self) -> dict:
        erros = []
        if self.erros_json:
            try:
                erros = json.loads(self.erros_json)
            except:
                pass
        
        return {
            'id': self.id,
            'organizacao_id': self.organizacao_id,
            'empresa_id': self.empresa_id,
            'nome_arquivo': self.nome_arquivo,
            'tipo_arquivo': self.tipo_arquivo,
            'tamanho_bytes': self.tamanho_bytes,
            'tamanho_formatado': self._format_size(self.tamanho_bytes),
            'status': self.status,
            'status_label': self._get_status_label(),
            'total_registros': self.total_registros,
            'registros_importados': self.registros_importados,
            'registros_duplicados': self.registros_duplicados,
            'registros_erro': self.registros_erro,
            'progresso': round((self.registros_importados / self.total_registros * 100) if self.total_registros > 0 else 0, 1),
            'erros': erros[:10],  # Primeiros 10 erros
            'total_erros': len(erros),
            'resumo': self.resumo,
            'iniciado_em': self.iniciado_em.isoformat() if self.iniciado_em else None,
            'finalizado_em': self.finalizado_em.isoformat() if self.finalizado_em else None,
            'created_at': self.created_at.isoformat() if self.created_at else None
        }
    
    def _get_status_label(self) -> str:
        labels = {
            'pending': 'Aguardando',
            'processing': 'Processando',
            'completed': 'Concluído',
            'failed': 'Falhou',
            'partial': 'Parcial'
        }
        return labels.get(self.status, self.status)
    
    def _format_size(self, size: int) -> str:
        if not size:
            return '-'
        for unit in ['B', 'KB', 'MB', 'GB']:
            if size < 1024:
                return f"{size:.1f} {unit}"
            size /= 1024
        return f"{size:.1f} TB"


class MapeamentoImportacao(Base):
    """Modelo de mapeamento salvo."""
    __tablename__ = 'mapeamentos_importacao'
    
    id = Column(Integer, primary_key=True, autoincrement=True)
    organizacao_id = Column(Integer, ForeignKey('organizacoes.id'))
    empresa_id = Column(Integer, ForeignKey('empresas.id'))
    contador_id = Column(Integer, ForeignKey('contadores.id'), nullable=False)
    
    nome = Column(String(100), nullable=False)
    descricao = Column(String(255))
    tipo_arquivo = Column(String(20), nullable=False)
    
    delimitador = Column(String(5), default=',')
    encoding = Column(String(20), default='utf-8')
    linha_cabecalho = Column(Integer, default=1)
    pular_linhas = Column(Integer, default=0)
    
    colunas_json = Column(Text, nullable=False)
    transformacoes_json = Column(Text)
    
    is_padrao = Column(Boolean, default=False)
    ativo = Column(Boolean, default=True)
    
    created_at = Column(DateTime, default=func.now())
    updated_at = Column(DateTime, default=func.now(), onupdate=func.now())
    
    def to_dict(self) -> dict:
        colunas = {}
        transformacoes = []
        try:
            colunas = json.loads(self.colunas_json) if self.colunas_json else {}
            transformacoes = json.loads(self.transformacoes_json) if self.transformacoes_json else []
        except:
            pass
        
        return {
            'id': self.id,
            'nome': self.nome,
            'descricao': self.descricao,
            'tipo_arquivo': self.tipo_arquivo,
            'delimitador': self.delimitador,
            'encoding': self.encoding,
            'linha_cabecalho': self.linha_cabecalho,
            'pular_linhas': self.pular_linhas,
            'colunas': colunas,
            'transformacoes': transformacoes,
            'is_padrao': self.is_padrao,
            'ativo': self.ativo,
            'created_at': self.created_at.isoformat() if self.created_at else None
        }


class RegistroImportado(Base):
    """Modelo de registro importado."""
    __tablename__ = 'registros_importados'
    
    id = Column(Integer, primary_key=True, autoincrement=True)
    importacao_id = Column(Integer, ForeignKey('importacoes.id'), nullable=False, index=True)
    empresa_id = Column(Integer, ForeignKey('empresas.id'), nullable=False)
    dados_mensal_id = Column(Integer, ForeignKey('dados_mensais.id'))
    
    linha_arquivo = Column(Integer)
    dados_originais_json = Column(Text)
    hash_registro = Column(String(64), index=True)
    
    status = Column(String(20), default='imported')
    mensagem_erro = Column(String(500))
    
    created_at = Column(DateTime, default=func.now())
    
    importacao = relationship("Importacao", back_populates="registros")


# =============================================================================
# DATA CLASSES
# =============================================================================

@dataclass
class DadoImportado:
    """Dados extraídos de uma linha."""
    ano: int
    mes: int
    receita: float = 0.0
    custos: float = 0.0
    despesas: float = 0.0
    impostos: float = 0.0
    folha: float = 0.0
    caixa: float = 0.0
    linha: int = 0
    dados_originais: Dict = field(default_factory=dict)
    erro: str = None
    
    @property
    def hash(self) -> str:
        """Gera hash único para detectar duplicatas."""
        conteudo = f"{self.ano}-{self.mes}-{self.receita}-{self.custos}-{self.despesas}"
        return hashlib.sha256(conteudo.encode()).hexdigest()[:32]


@dataclass
class ResultadoImportacao:
    """Resultado de uma importação."""
    sucesso: bool
    total: int = 0
    importados: int = 0
    duplicados: int = 0
    erros: int = 0
    dados: List[DadoImportado] = field(default_factory=list)
    mensagens_erro: List[Dict] = field(default_factory=list)
    preview: List[Dict] = field(default_factory=list)


# =============================================================================
# PARSERS
# =============================================================================

def parse_valor_monetario(valor: Any) -> float:
    """Converte valor monetário para float."""
    if valor is None:
        return 0.0
    
    if isinstance(valor, (int, float)):
        return float(valor)
    
    texto = str(valor).strip()
    if not texto or texto.lower() in ('', '-', 'null', 'none', 'nan'):
        return 0.0
    
    # Remove símbolos de moeda
    texto = re.sub(r'[R$\s]', '', texto)
    
    # Detecta formato brasileiro (1.234,56) vs americano (1,234.56)
    if ',' in texto and '.' in texto:
        if texto.rfind(',') > texto.rfind('.'):
            # Formato brasileiro: 1.234,56
            texto = texto.replace('.', '').replace(',', '.')
        else:
            # Formato americano: 1,234.56
            texto = texto.replace(',', '')
    elif ',' in texto:
        # Pode ser 1234,56 (brasileiro) ou 1,234 (americano)
        partes = texto.split(',')
        if len(partes) == 2 and len(partes[1]) <= 2:
            # Provavelmente brasileiro
            texto = texto.replace(',', '.')
        else:
            # Provavelmente americano
            texto = texto.replace(',', '')
    
    try:
        return float(texto)
    except (ValueError, InvalidOperation):
        return 0.0


def parse_data(valor: Any, formatos: List[str] = None) -> Optional[date]:
    """Converte valor para data."""
    if valor is None:
        return None
    
    if isinstance(valor, datetime):
        return valor.date()
    
    if isinstance(valor, date):
        return valor
    
    texto = str(valor).strip()
    if not texto:
        return None
    
    formatos = formatos or [
        '%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%Y/%m/%d',
        '%d/%m/%y', '%d-%m-%y', '%m/%d/%Y', '%m-%d-%Y',
        '%d.%m.%Y', '%Y.%m.%d'
    ]
    
    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    
    # Tenta extrair ano e mês de formatos como "jan/2024", "2024-01", etc.
    meses = {
        'jan': 1, 'fev': 2, 'mar': 3, 'abr': 4, 'mai': 5, 'jun': 6,
        'jul': 7, 'ago': 8, 'set': 9, 'out': 10, 'nov': 11, 'dez': 12,
        'january': 1, 'february': 2, 'march': 3, 'april': 4, 'may': 5, 'june': 6,
        'july': 7, 'august': 8, 'september': 9, 'october': 10, 'november': 11, 'december': 12
    }
    
    texto_lower = texto.lower()
    for nome, num in meses.items():
        if nome in texto_lower:
            # Extrai ano
            anos = re.findall(r'20\d{2}|19\d{2}', texto)
            if anos:
                return date(int(anos[0]), num, 1)
    
    # Tenta formato YYYY-MM
    match = re.match(r'(\d{4})[/-](\d{1,2})', texto)
    if match:
        return date(int(match.group(1)), int(match.group(2)), 1)
    
    return None


def calcular_hash_arquivo(conteudo: bytes) -> str:
    """Calcula hash SHA256 do arquivo."""
    return hashlib.sha256(conteudo).hexdigest()


# =============================================================================
# IMPORTADORES
# =============================================================================

def importar_csv(
    conteudo: bytes,
    mapeamento: Dict,
    encoding: str = 'utf-8',
    delimitador: str = ',',
    pular_linhas: int = 0
) -> ResultadoImportacao:
    """Importa dados de arquivo CSV."""
    resultado = ResultadoImportacao(sucesso=True)
    
    try:
        texto = conteudo.decode(encoding)
    except UnicodeDecodeError:
        # Tenta outros encodings
        for enc in ['latin-1', 'cp1252', 'iso-8859-1']:
            try:
                texto = conteudo.decode(enc)
                break
            except:
                continue
        else:
            resultado.sucesso = False
            resultado.mensagens_erro.append({'linha': 0, 'erro': 'Encoding não suportado'})
            return resultado
    
    linhas = texto.strip().split('\n')
    
    # Pula linhas iniciais
    linhas = linhas[pular_linhas:]
    
    if not linhas:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': 'Arquivo vazio'})
        return resultado
    
    # Primeira linha é cabeçalho
    cabecalho = [c.strip().strip('"').lower() for c in linhas[0].split(delimitador)]
    
    # Processa linhas de dados
    for i, linha in enumerate(linhas[1:], start=2 + pular_linhas):
        if not linha.strip():
            continue
        
        resultado.total += 1
        valores = [v.strip().strip('"') for v in linha.split(delimitador)]
        
        # Cria dicionário com dados originais
        dados_orig = dict(zip(cabecalho, valores))
        
        try:
            dado = extrair_dados_linha(dados_orig, mapeamento, i)
            if dado.erro:
                resultado.erros += 1
                resultado.mensagens_erro.append({'linha': i, 'erro': dado.erro, 'dados': dados_orig})
            else:
                resultado.dados.append(dado)
                resultado.importados += 1
        except Exception as e:
            resultado.erros += 1
            resultado.mensagens_erro.append({'linha': i, 'erro': str(e), 'dados': dados_orig})
    
    # Preview (primeiras 5 linhas)
    resultado.preview = [
        dict(zip(cabecalho, [v.strip().strip('"') for v in l.split(delimitador)]))
        for l in linhas[1:6] if l.strip()
    ]
    
    return resultado


def importar_excel(
    conteudo: bytes,
    mapeamento: Dict,
    sheet: str = None,
    linha_cabecalho: int = 1,
    pular_linhas: int = 0
) -> ResultadoImportacao:
    """Importa dados de arquivo Excel."""
    resultado = ResultadoImportacao(sucesso=True)
    
    if not EXCEL_AVAILABLE:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': 'Biblioteca openpyxl não disponível'})
        return resultado
    
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        
        # Lê todas as linhas
        linhas = list(ws.iter_rows(values_only=True))
        
        if len(linhas) < linha_cabecalho:
            resultado.sucesso = False
            resultado.mensagens_erro.append({'linha': 0, 'erro': 'Arquivo vazio ou cabeçalho não encontrado'})
            return resultado
        
        # Cabeçalho
        cabecalho = [str(c).strip().lower() if c else f'col_{i}' for i, c in enumerate(linhas[linha_cabecalho - 1])]
        
        # Processa linhas de dados
        for i, linha in enumerate(linhas[linha_cabecalho + pular_linhas:], start=linha_cabecalho + pular_linhas + 1):
            if not any(linha):
                continue
            
            resultado.total += 1
            dados_orig = dict(zip(cabecalho, linha))
            
            try:
                dado = extrair_dados_linha(dados_orig, mapeamento, i)
                if dado.erro:
                    resultado.erros += 1
                    resultado.mensagens_erro.append({'linha': i, 'erro': dado.erro, 'dados': str(dados_orig)})
                else:
                    resultado.dados.append(dado)
                    resultado.importados += 1
            except Exception as e:
                resultado.erros += 1
                resultado.mensagens_erro.append({'linha': i, 'erro': str(e), 'dados': str(dados_orig)})
        
        # Preview
        for linha in linhas[linha_cabecalho:linha_cabecalho + 5]:
            if any(linha):
                resultado.preview.append(dict(zip(cabecalho, [str(v) if v else '' for v in linha])))
        
        wb.close()
        
    except Exception as e:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': f'Erro ao ler Excel: {str(e)}'})
    
    return resultado


def importar_ofx(conteudo: bytes) -> ResultadoImportacao:
    """Importa dados de arquivo OFX (extrato bancário)."""
    resultado = ResultadoImportacao(sucesso=True)
    
    try:
        texto = conteudo.decode('utf-8', errors='ignore')
    except:
        texto = conteudo.decode('latin-1', errors='ignore')
    
    # Parser simples de OFX
    transacoes = []
    
    # Extrai transações usando regex
    stmttrn_pattern = r'<STMTTRN>(.*?)</STMTTRN>'
    matches = re.findall(stmttrn_pattern, texto, re.DOTALL | re.IGNORECASE)
    
    for i, match in enumerate(matches, start=1):
        resultado.total += 1
        
        # Extrai campos
        def get_field(name):
            pattern = f'<{name}>([^<\n]+)'
            m = re.search(pattern, match, re.IGNORECASE)
            return m.group(1).strip() if m else None
        
        try:
            dtposted = get_field('DTPOSTED')
            trnamt = get_field('TRNAMT')
            memo = get_field('MEMO') or get_field('NAME') or ''
            
            if dtposted and trnamt:
                # Parse data (formato YYYYMMDD ou YYYYMMDDHHMMSS)
                ano = int(dtposted[:4])
                mes = int(dtposted[4:6])
                valor = parse_valor_monetario(trnamt)
                
                dado = DadoImportado(
                    ano=ano,
                    mes=mes,
                    receita=valor if valor > 0 else 0,
                    despesas=abs(valor) if valor < 0 else 0,
                    linha=i,
                    dados_originais={'data': dtposted, 'valor': trnamt, 'descricao': memo}
                )
                resultado.dados.append(dado)
                resultado.importados += 1
                
                if len(resultado.preview) < 5:
                    resultado.preview.append({
                        'data': f"{dtposted[:4]}-{dtposted[4:6]}-{dtposted[6:8]}",
                        'valor': trnamt,
                        'descricao': memo
                    })
        except Exception as e:
            resultado.erros += 1
            resultado.mensagens_erro.append({'linha': i, 'erro': str(e)})
    
    if not matches:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': 'Nenhuma transação encontrada no arquivo OFX'})
    
    return resultado


def importar_xml_nfe(conteudo: bytes) -> ResultadoImportacao:
    """Importa dados de arquivo XML de NFe."""
    resultado = ResultadoImportacao(sucesso=True)
    
    try:
        root = ET.fromstring(conteudo)
    except ET.ParseError as e:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': f'XML inválido: {str(e)}'})
        return resultado
    
    # Namespaces comuns de NFe
    namespaces = {
        'nfe': 'http://www.portalfiscal.inf.br/nfe',
        '': 'http://www.portalfiscal.inf.br/nfe'
    }
    
    # Tenta encontrar elementos com ou sem namespace
    def find_element(parent, tag):
        for ns_prefix, ns_uri in namespaces.items():
            el = parent.find(f'{{{ns_uri}}}{tag}')
            if el is not None:
                return el
        return parent.find(tag)
    
    def find_all(parent, tag):
        for ns_prefix, ns_uri in namespaces.items():
            els = parent.findall(f'.//{{{ns_uri}}}{tag}')
            if els:
                return els
        return parent.findall(f'.//{tag}')
    
    # Procura NFes no documento
    nfes = find_all(root, 'NFe') or find_all(root, 'infNFe') or [root]
    
    for i, nfe in enumerate(nfes, start=1):
        resultado.total += 1
        
        try:
            # Busca informações da NFe
            ide = find_element(nfe, 'ide')
            total = find_element(nfe, 'total')
            
            # Data de emissão
            dhemi = None
            if ide is not None:
                dhemi_el = find_element(ide, 'dhEmi') or find_element(ide, 'dEmi')
                if dhemi_el is not None and dhemi_el.text:
                    dhemi = dhemi_el.text
            
            # Valores
            icms_tot = find_element(total, 'ICMSTot') if total is not None else None
            
            vnf = 0  # Valor total NF
            vprod = 0  # Valor produtos
            vipi = 0  # IPI
            vicms = 0  # ICMS
            
            if icms_tot is not None:
                vnf_el = find_element(icms_tot, 'vNF')
                vprod_el = find_element(icms_tot, 'vProd')
                vipi_el = find_element(icms_tot, 'vIPI')
                vicms_el = find_element(icms_tot, 'vICMS')
                
                vnf = parse_valor_monetario(vnf_el.text if vnf_el is not None else 0)
                vprod = parse_valor_monetario(vprod_el.text if vprod_el is not None else 0)
                vipi = parse_valor_monetario(vipi_el.text if vipi_el is not None else 0)
                vicms = parse_valor_monetario(vicms_el.text if vicms_el is not None else 0)
            
            # Parse data
            if dhemi:
                data = parse_data(dhemi[:10])
                if data:
                    dado = DadoImportado(
                        ano=data.year,
                        mes=data.month,
                        receita=vnf,
                        custos=vprod - vnf if vprod > vnf else 0,
                        impostos=vipi + vicms,
                        linha=i,
                        dados_originais={
                            'data_emissao': dhemi,
                            'valor_nf': vnf,
                            'valor_produtos': vprod,
                            'ipi': vipi,
                            'icms': vicms
                        }
                    )
                    resultado.dados.append(dado)
                    resultado.importados += 1
                    
                    if len(resultado.preview) < 5:
                        resultado.preview.append(dado.dados_originais)
                else:
                    resultado.erros += 1
                    resultado.mensagens_erro.append({'linha': i, 'erro': f'Data inválida: {dhemi}'})
            else:
                resultado.erros += 1
                resultado.mensagens_erro.append({'linha': i, 'erro': 'Data de emissão não encontrada'})
                
        except Exception as e:
            resultado.erros += 1
            resultado.mensagens_erro.append({'linha': i, 'erro': str(e)})
    
    if not resultado.dados:
        resultado.sucesso = False
        resultado.mensagens_erro.append({'linha': 0, 'erro': 'Nenhuma NFe válida encontrada'})
    
    return resultado


def extrair_dados_linha(dados: Dict, mapeamento: Dict, linha: int) -> DadoImportado:
    """Extrai dados de uma linha usando o mapeamento."""
    
    def get_valor(campo: str) -> Any:
        config = mapeamento.get(campo, {})
        if isinstance(config, str):
            coluna = config
        else:
            coluna = config.get('coluna', campo)
        
        # Busca pela coluna (case insensitive)
        for k, v in dados.items():
            if k.lower() == coluna.lower():
                return v
        return None
    
    # Extrai data
    data_valor = get_valor('data') or get_valor('competencia') or get_valor('periodo')
    data = parse_data(data_valor)
    
    if not data:
        # Tenta extrair ano e mês separadamente
        ano = get_valor('ano')
        mes = get_valor('mes')
        if ano and mes:
            try:
                data = date(int(ano), int(mes), 1)
            except:
                pass
    
    if not data:
        return DadoImportado(
            ano=0, mes=0, linha=linha,
            dados_originais=dados,
            erro=f"Data inválida ou não encontrada: {data_valor}"
        )
    
    return DadoImportado(
        ano=data.year,
        mes=data.month,
        receita=parse_valor_monetario(get_valor('receita') or get_valor('receita_bruta') or get_valor('faturamento')),
        custos=parse_valor_monetario(get_valor('custos') or get_valor('custo') or get_valor('cmv')),
        despesas=parse_valor_monetario(get_valor('despesas') or get_valor('despesas_operacionais')),
        impostos=parse_valor_monetario(get_valor('impostos') or get_valor('tributos')),
        folha=parse_valor_monetario(get_valor('folha') or get_valor('folha_pagamento') or get_valor('salarios')),
        caixa=parse_valor_monetario(get_valor('caixa') or get_valor('saldo_caixa') or get_valor('saldo')),
        linha=linha,
        dados_originais=dados
    )


# =============================================================================
# FUNÇÕES DE SERVIÇO
# =============================================================================

def detectar_tipo_arquivo(nome_arquivo: str, conteudo: bytes = None) -> str:
    """Detecta tipo de arquivo."""
    nome = nome_arquivo.lower()
    
    if nome.endswith('.xlsx') or nome.endswith('.xls'):
        return 'xlsx'
    elif nome.endswith('.ofx') or nome.endswith('.qif'):
        return 'ofx'
    elif nome.endswith('.xml'):
        # Verifica se é NFe
        if conteudo:
            texto = conteudo[:2000].decode('utf-8', errors='ignore').lower()
            if 'nfe' in texto or 'portalfiscal' in texto:
                return 'xml_nfe'
        return 'xml'
    else:
        return 'csv'


def detectar_mapeamento(conteudo: bytes, tipo: str) -> Dict:
    """Tenta detectar mapeamento automático baseado nos cabeçalhos."""
    mapeamento = {}
    
    # Mapeamentos conhecidos (coluna no arquivo -> campo no sistema)
    conhecidos = {
        'data': ['data', 'date', 'competencia', 'periodo', 'mes_ano', 'ref', 'referencia'],
        'ano': ['ano', 'year'],
        'mes': ['mes', 'month'],
        'receita': ['receita', 'receita_bruta', 'faturamento', 'vendas', 'revenue', 'receitas', 'total_vendas'],
        'custos': ['custos', 'custo', 'cmv', 'cpv', 'costs', 'custo_produtos'],
        'despesas': ['despesas', 'despesas_operacionais', 'expenses', 'gastos', 'despesa'],
        'impostos': ['impostos', 'tributos', 'taxes', 'imposto'],
        'folha': ['folha', 'folha_pagamento', 'salarios', 'payroll', 'pessoal'],
        'caixa': ['caixa', 'saldo_caixa', 'saldo', 'cash', 'disponivel', 'banco']
    }
    
    # Tenta extrair cabeçalho
    cabecalho = []
    
    if tipo == 'csv':
        try:
            texto = conteudo.decode('utf-8', errors='ignore')
            primeira_linha = texto.split('\n')[0]
            cabecalho = [c.strip().strip('"').lower() for c in primeira_linha.split(',')]
        except:
            pass
    
    elif tipo == 'xlsx' and EXCEL_AVAILABLE:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
            ws = wb.active
            primeira_linha = next(ws.iter_rows(max_row=1, values_only=True))
            cabecalho = [str(c).strip().lower() if c else '' for c in primeira_linha]
            wb.close()
        except:
            pass
    
    # Mapeia colunas encontradas
    for campo, aliases in conhecidos.items():
        for col in cabecalho:
            if col in aliases or any(alias in col for alias in aliases):
                mapeamento[campo] = col
                break
    
    return mapeamento


def processar_importacao(
    conteudo: bytes,
    nome_arquivo: str,
    empresa_id: int,
    contador_id: int,
    mapeamento: Dict = None,
    organizacao_id: int = None
) -> Dict:
    """Processa uma importação completa."""
    
    tipo = detectar_tipo_arquivo(nome_arquivo, conteudo)
    hash_arquivo = calcular_hash_arquivo(conteudo)
    
    with get_db() as db:
        # Verifica duplicata de arquivo
        existente = db.query(Importacao).filter(
            Importacao.empresa_id == empresa_id,
            Importacao.hash_arquivo == hash_arquivo,
            Importacao.status == 'completed'
        ).first()
        
        if existente:
            return {
                'sucesso': False,
                'erro': 'Este arquivo já foi importado anteriormente',
                'importacao_anterior': existente.to_dict()
            }
        
        # Cria registro de importação
        importacao = Importacao(
            organizacao_id=organizacao_id,
            empresa_id=empresa_id,
            contador_id=contador_id,
            nome_arquivo=nome_arquivo,
            tipo_arquivo=tipo,
            tamanho_bytes=len(conteudo),
            hash_arquivo=hash_arquivo,
            status='processing',
            iniciado_em=datetime.now()
        )
        db.add(importacao)
        db.flush()
        
        # Detecta mapeamento se não fornecido
        if not mapeamento:
            mapeamento = detectar_mapeamento(conteudo, tipo)
        
        importacao.mapeamento_json = json.dumps(mapeamento)
        
        # Processa conforme tipo
        if tipo == 'xlsx':
            resultado = importar_excel(conteudo, mapeamento)
        elif tipo == 'ofx':
            resultado = importar_ofx(conteudo)
        elif tipo == 'xml_nfe':
            resultado = importar_xml_nfe(conteudo)
        else:
            resultado = importar_csv(conteudo, mapeamento)
        
        importacao.total_registros = resultado.total
        
        # Salva dados
        from data.database import salvar_dados_mensais
        
        for dado in resultado.dados:
            # Verifica duplicata por competência
            hash_reg = dado.hash
            
            existente_reg = db.query(RegistroImportado).filter(
                RegistroImportado.empresa_id == empresa_id,
                RegistroImportado.hash_registro == hash_reg
            ).first()
            
            if existente_reg:
                importacao.registros_duplicados += 1
                registro = RegistroImportado(
                    importacao_id=importacao.id,
                    empresa_id=empresa_id,
                    linha_arquivo=dado.linha,
                    dados_originais_json=json.dumps(dado.dados_originais),
                    hash_registro=hash_reg,
                    status='duplicate',
                    mensagem_erro='Registro duplicado'
                )
            else:
                # Salva dados mensais
                try:
                    salvar_dados_mensais(empresa_id, {
                        'ano': dado.ano,
                        'mes': dado.mes,
                        'receita': dado.receita,
                        'custos': dado.custos,
                        'despesas': dado.despesas,
                        'impostos': dado.impostos,
                        'folha': dado.folha,
                        'caixa': dado.caixa
                    })
                    
                    registro = RegistroImportado(
                        importacao_id=importacao.id,
                        empresa_id=empresa_id,
                        linha_arquivo=dado.linha,
                        dados_originais_json=json.dumps(dado.dados_originais),
                        hash_registro=hash_reg,
                        status='imported'
                    )
                    importacao.registros_importados += 1
                    
                except Exception as e:
                    registro = RegistroImportado(
                        importacao_id=importacao.id,
                        empresa_id=empresa_id,
                        linha_arquivo=dado.linha,
                        dados_originais_json=json.dumps(dado.dados_originais),
                        hash_registro=hash_reg,
                        status='error',
                        mensagem_erro=str(e)
                    )
                    importacao.registros_erro += 1
            
            db.add(registro)
        
        # Registra erros de parsing
        for erro in resultado.mensagens_erro:
            importacao.registros_erro += 1
        
        importacao.erros_json = json.dumps(resultado.mensagens_erro)
        
        # Define status final
        if importacao.registros_importados == 0:
            importacao.status = 'failed'
        elif importacao.registros_erro > 0 or importacao.registros_duplicados > 0:
            importacao.status = 'partial'
        else:
            importacao.status = 'completed'
        
        importacao.finalizado_em = datetime.now()
        importacao.resumo = f"Importados: {importacao.registros_importados}, Duplicados: {importacao.registros_duplicados}, Erros: {importacao.registros_erro}"
        
        return {
            'sucesso': importacao.status in ('completed', 'partial'),
            'importacao': importacao.to_dict(),
            'preview': resultado.preview
        }


def preview_importacao(
    conteudo: bytes,
    nome_arquivo: str,
    mapeamento: Dict = None
) -> Dict:
    """Gera preview de importação sem salvar."""
    
    tipo = detectar_tipo_arquivo(nome_arquivo, conteudo)
    
    if not mapeamento:
        mapeamento = detectar_mapeamento(conteudo, tipo)
    
    if tipo == 'xlsx':
        resultado = importar_excel(conteudo, mapeamento)
    elif tipo == 'ofx':
        resultado = importar_ofx(conteudo)
    elif tipo == 'xml_nfe':
        resultado = importar_xml_nfe(conteudo)
    else:
        resultado = importar_csv(conteudo, mapeamento)
    
    return {
        'tipo_detectado': tipo,
        'mapeamento_sugerido': mapeamento,
        'total_registros': resultado.total,
        'registros_validos': resultado.importados,
        'registros_erro': resultado.erros,
        'preview': resultado.preview,
        'erros': resultado.mensagens_erro[:5],
        'dados_extraidos': [
            {
                'linha': d.linha,
                'ano': d.ano,
                'mes': d.mes,
                'receita': d.receita,
                'custos': d.custos,
                'despesas': d.despesas,
                'impostos': d.impostos,
                'folha': d.folha,
                'caixa': d.caixa
            }
            for d in resultado.dados[:10]
        ]
    }


def listar_importacoes(empresa_id: int = None, organizacao_id: int = None, limite: int = 20) -> List[Dict]:
    """Lista importações."""
    with get_db() as db:
        query = db.query(Importacao)
        
        if empresa_id:
            query = query.filter(Importacao.empresa_id == empresa_id)
        elif organizacao_id:
            query = query.filter(Importacao.organizacao_id == organizacao_id)
        
        importacoes = query.order_by(desc(Importacao.created_at)).limit(limite).all()
        return [i.to_dict() for i in importacoes]


def obter_importacao(importacao_id: int) -> Optional[Dict]:
    """Obtém detalhes de uma importação."""
    with get_db() as db:
        importacao = db.query(Importacao).filter(Importacao.id == importacao_id).first()
        return importacao.to_dict() if importacao else None


# Mapeamentos
def salvar_mapeamento(
    nome: str,
    tipo_arquivo: str,
    colunas: Dict,
    contador_id: int,
    empresa_id: int = None,
    organizacao_id: int = None,
    **kwargs
) -> Dict:
    """Salva um mapeamento."""
    with get_db() as db:
        mapeamento = MapeamentoImportacao(
            nome=nome,
            tipo_arquivo=tipo_arquivo,
            colunas_json=json.dumps(colunas),
            contador_id=contador_id,
            empresa_id=empresa_id,
            organizacao_id=organizacao_id,
            delimitador=kwargs.get('delimitador', ','),
            encoding=kwargs.get('encoding', 'utf-8'),
            linha_cabecalho=kwargs.get('linha_cabecalho', 1),
            pular_linhas=kwargs.get('pular_linhas', 0),
            transformacoes_json=json.dumps(kwargs.get('transformacoes', []))
        )
        db.add(mapeamento)
        db.flush()
        return mapeamento.to_dict()


def listar_mapeamentos(empresa_id: int = None, organizacao_id: int = None) -> List[Dict]:
    """Lista mapeamentos salvos."""
    with get_db() as db:
        query = db.query(MapeamentoImportacao).filter(MapeamentoImportacao.ativo == True)
        
        if empresa_id:
            query = query.filter(
                (MapeamentoImportacao.empresa_id == empresa_id) |
                (MapeamentoImportacao.empresa_id.is_(None))
            )
        elif organizacao_id:
            query = query.filter(
                (MapeamentoImportacao.organizacao_id == organizacao_id) |
                (MapeamentoImportacao.organizacao_id.is_(None))
            )
        
        mapeamentos = query.order_by(desc(MapeamentoImportacao.is_padrao), MapeamentoImportacao.nome).all()
        return [m.to_dict() for m in mapeamentos]
