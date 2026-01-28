#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importador Excel (XLSX/XLS)
Suporta tanto XLSX modernos quanto XLS antigos (Excel 97-2003)
"""

import io
import struct
from typing import List, Dict, Tuple

from .base import ImportadorBase

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False


class ImportadorExcel(ImportadorBase):
    """Importador para arquivos Excel (XLS e XLSX)."""
    
    TIPO = "xlsx"
    EXTENSOES = ['.xlsx', '.xls', '.xlsm']
    
    def ler_arquivo(self, conteudo: bytes, nome_arquivo: str) -> Tuple[List[Dict], List[str]]:
        """Lê arquivo Excel detectando automaticamente o formato."""
        
        # Detectar tipo de arquivo pelo magic number
        if conteudo[:4] == b'\xd0\xcf\x11\xe0':
            # XLS antigo (formato OLE/BIFF)
            return self._ler_xls_antigo(conteudo, nome_arquivo)
        elif conteudo[:2] == b'PK':
            # XLSX moderno (é um arquivo ZIP)
            return self._ler_xlsx_moderno(conteudo, nome_arquivo)
        else:
            # Tentar como XLSX
            try:
                return self._ler_xlsx_moderno(conteudo, nome_arquivo)
            except:
                return self._ler_xls_antigo(conteudo, nome_arquivo)
    
    def _ler_xlsx_moderno(self, conteudo: bytes, nome_arquivo: str) -> Tuple[List[Dict], List[str]]:
        """Lê arquivo XLSX usando openpyxl."""
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl não instalado. Use: pip install openpyxl")
        
        dados = []
        colunas = []
        
        # Abre workbook
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        ws = wb.active
        
        # Encontra primeira linha com dados (header)
        header_row = 1
        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=10), 1):
            valores = [cell.value for cell in row if cell.value]
            if len(valores) >= 2:
                header_row = row_num
                break
        
        # Lê headers
        headers = []
        for cell in ws[header_row]:
            valor = cell.value
            if valor:
                headers.append(str(valor).strip())
            else:
                headers.append(f"Coluna_{cell.column}")
        
        colunas = [h for h in headers if not h.startswith('Coluna_') or headers.count(h) == 1]
        
        # Lê dados
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(row):
                continue
            
            row_dict = {}
            for i, valor in enumerate(row):
                if i < len(headers):
                    header = headers[i]
                    if valor is None:
                        row_dict[header] = ''
                    elif hasattr(valor, 'strftime'):
                        row_dict[header] = valor.strftime('%Y-%m-%d')
                    else:
                        row_dict[header] = valor
            
            if any(v for v in row_dict.values() if v not in [None, '', 0]):
                dados.append(row_dict)
        
        wb.close()
        return dados, colunas
    
    def _ler_xls_antigo(self, conteudo: bytes, nome_arquivo: str) -> Tuple[List[Dict], List[str]]:
        """Lê arquivo XLS antigo usando xlrd ou parser binário."""
        
        # Tentar usar xlrd primeiro
        if XLRD_AVAILABLE:
            try:
                return self._ler_xls_com_xlrd(conteudo, nome_arquivo)
            except:
                pass
        
        # Fallback: usar parser binário
        return self._ler_xls_binario(conteudo, nome_arquivo)
    
    def _ler_xls_com_xlrd(self, conteudo: bytes, nome_arquivo: str) -> Tuple[List[Dict], List[str]]:
        """Lê XLS usando xlrd."""
        import xlrd
        
        wb = xlrd.open_workbook(file_contents=conteudo)
        ws = wb.sheet_by_index(0)
        
        dados = []
        
        # Encontrar header
        header_row = 0
        for row_idx in range(min(10, ws.nrows)):
            row = [ws.cell_value(row_idx, col) for col in range(ws.ncols)]
            valores = [v for v in row if v]
            if len(valores) >= 2:
                header_row = row_idx
                break
        
        # Headers
        headers = []
        for col in range(ws.ncols):
            val = ws.cell_value(header_row, col)
            if val:
                headers.append(str(val).strip())
            else:
                headers.append(f"Coluna_{col+1}")
        
        colunas = [h for h in headers if not h.startswith('Coluna_')]
        
        # Dados
        for row_idx in range(header_row + 1, ws.nrows):
            row = [ws.cell_value(row_idx, col) for col in range(ws.ncols)]
            
            if not any(row):
                continue
            
            row_dict = {}
            for i, valor in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = valor if valor else ''
            
            if any(v for v in row_dict.values() if v not in [None, '', 0]):
                dados.append(row_dict)
        
        return dados, colunas
    
    def _ler_xls_binario(self, conteudo: bytes, nome_arquivo: str) -> Tuple[List[Dict], List[str]]:
        """
        Parser binário para XLS antigos quando xlrd não está disponível.
        Extrai textos e valores do arquivo binário.
        """
        import re
        
        dados = []
        colunas = []
        
        # Extrair strings UTF-16LE
        strings = self._extrair_strings_utf16(conteudo)
        
        # Extrair valores numéricos
        valores = self._extrair_valores_ieee754(conteudo)
        
        # Tentar reconstruir estrutura básica
        # Para balancetes, os dados são: código, descrição, saldo_ant, débito, crédito, saldo_atual
        
        # Colunas comuns de balancete
        colunas = ['Código', 'Descrição', 'Saldo Anterior', 'Débito', 'Crédito', 'Saldo Atual']
        
        # Os dados extraídos são usados pelo importador de balancetes
        # Aqui apenas retornamos os textos encontrados
        for i, s in enumerate(strings[:50]):
            dados.append({
                'linha': i,
                'texto': s,
                'tipo': 'string'
            })
        
        for i, v in enumerate(valores[:50]):
            dados.append({
                'linha': 100 + i,
                'valor': v,
                'tipo': 'numero'
            })
        
        return dados, colunas
    
    def _extrair_strings_utf16(self, data: bytes) -> List[str]:
        """Extrai strings UTF-16LE do binário."""
        import re
        
        strings = []
        i = 0
        
        while i < len(data) - 1:
            if data[i] >= 32 and data[i] < 127 and data[i+1] == 0:
                texto = []
                j = i
                while j < len(data) - 1:
                    low = data[j]
                    high = data[j+1] if j+1 < len(data) else 0
                    char_code = low | (high << 8)
                    
                    if (32 <= char_code < 127) or (0xC0 <= char_code <= 0xFF):
                        texto.append(chr(char_code))
                        j += 2
                    elif char_code == 0:
                        break
                    else:
                        break
                
                if len(texto) >= 3:
                    s = ''.join(texto).strip()
                    if s:
                        strings.append(s)
                    i = j
                else:
                    i += 2
            else:
                i += 1
        
        # Remover duplicatas
        seen = set()
        unicos = []
        for s in strings:
            if s not in seen:
                seen.add(s)
                unicos.append(s)
        
        return unicos
    
    def _extrair_valores_ieee754(self, data: bytes) -> List[float]:
        """Extrai valores numéricos IEEE 754 double."""
        valores = set()
        
        for i in range(0, len(data) - 8):
            try:
                val = struct.unpack('<d', data[i:i+8])[0]
                if 0.01 <= abs(val) <= 1e11:
                    rounded = round(val, 2)
                    if abs(val - rounded) < 0.001:
                        valores.add(rounded)
            except:
                pass
        
        return sorted(valores, reverse=True)
